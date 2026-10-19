import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (4).pptx"
prs = pptx.Presentation(pptx_path)

print(f"=== PPTX: {pptx_path} (Total Slides: {len(prs.slides)}) ===")
for idx, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {idx+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip().replace('\n', ' | ')
            if txt:
                print(f"  Shape {s_idx}: {txt}")
