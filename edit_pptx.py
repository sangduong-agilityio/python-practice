import sys
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (3).pptx"
backup_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (3)_backup.pptx"

# Make backup
shutil.copyfile(pptx_path, backup_path)
print(f"Backup created at: {backup_path}")

prs = pptx.Presentation(pptx_path)

def add_header(slide, title_text, subtitle_text):
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(9.0), Inches(0.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42) # dark navy

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.85), Inches(9.0), Inches(0.25))
    tf2 = sub_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle_text.upper()
    p2.font.name = 'Calibri'
    p2.font.size = Pt(9)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(100, 116, 139) # slate gray

def add_card(slide, left, top, width, height, title, items, bg_rgb=RGBColor(241, 245, 249)):
    # Shape container
    shape = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_rgb
    shape.line.color.rgb = RGBColor(226, 232, 240)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_bottom = Inches(0.15)

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.name = 'Calibri'
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(30, 58, 138)

    for item in items:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.name = 'Calibri'
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(4)

print("Script template ready.")
