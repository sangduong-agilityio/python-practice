import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

for s_idx in [3, 7, 8]: # Slides 4, 8, 9
    slide = prs.slides[s_idx]
    print(f"\n================ SLIDE {s_idx+1} ================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.color and r.font.color.rgb:
                        print(f"Text: '{r.text[:30]}' Color: {r.font.color.rgb}")
