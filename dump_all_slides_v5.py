import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (5).pptx"
prs = pptx.Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")

for s_idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {s_idx+1} ====================")
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                t = p.text.strip()
                if t:
                    print(f"  [S{s_idx+1} P{p_idx}]: {t}")
        if shape.has_table:
            print(f"  [TABLE {shape.name}]: {len(shape.table.rows)} rows x {len(shape.table.columns)} cols")
