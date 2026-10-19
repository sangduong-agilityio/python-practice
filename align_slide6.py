import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

# Palette
C_WHITE = RGBColor(255, 255, 255)
C_SLATE_100 = RGBColor(241, 245, 249)
C_SLATE_200 = RGBColor(226, 232, 240)
C_SLATE_300 = RGBColor(203, 213, 225)
C_SLATE_400 = RGBColor(148, 163, 184)
C_SLATE_700 = RGBColor(51, 65, 85)
C_SLATE_800 = RGBColor(30, 41, 59)
C_SLATE_900 = RGBColor(15, 23, 42)
C_NAVY_CARD = RGBColor(19, 30, 50)
C_NAVY_DARK = RGBColor(11, 18, 32)

C_CYAN = RGBColor(56, 189, 248) # #38BDF8
C_CYAN_DARK = RGBColor(2, 132, 199)
C_CORAL_RED = RGBColor(248, 113, 113) # #F87171
C_RED_VIBRANT = RGBColor(239, 68, 68)
C_RED_HEADER = RGBColor(69, 26, 26)
C_RED_CELL = RGBColor(32, 20, 26)
C_RED_TEXT = RGBColor(254, 202, 202)

C_MINT_GREEN = RGBColor(52, 211, 153) # #34D399
C_GREEN_HEADER = RGBColor(6, 78, 59)
C_GREEN_CELL = RGBColor(8, 38, 32)
C_GREEN_TEXT = RGBColor(167, 243, 208)

slide6 = prs.slides[5]

# Clean up Slide 6: Delete all shapes except slide number (Shape 0), Title (Shape 1), Subtitle (Shape 2)
# We will rebuild the 3 columns with 100% mathematically exact, pixel-perfect alignment!

# Let's inspect slide6 shapes and remove old column elements
shapes_to_remove = []
for idx in range(len(slide6.shapes) - 1, -1, -1):
    shape = slide6.shapes[idx]
    if shape.name not in ['Google Shape;313;p8', 'TextBox 313', 'TextBox 314']:
        shapes_to_remove.append(shape)

for shape in shapes_to_remove:
    sp = shape._element
    sp.getparent().remove(sp)

print(f"Remaining shapes on Slide 6: {len(slide6.shapes)}")

# Now let's rebuild the 3 symmetrical columns on Slide 6!
col_width = 2670048
gap = 201168
left_margin = 365760
card_top = 960000
card_height = 3680000

cols_data = [
    {
        "badge": "1NF",
        "title": "First Normal Form",
        "rule": "Each cell must hold exactly one atomic value (no comma-separated lists/arrays).",
        "before_headers": ["id", "title", "assignee_emails"],
        "before_rows": [
            ["1", "Fix login", "tuyen@gmail.com, huy@gmail.com"]
        ],
        "highlight_col": 2,
        "col_widths": [380000, 780000, 1254016],
        "after_box1_title": "tasks",
        "after_box1_sub": "task_id (PK), title",
        "after_box2_title": "task_assignees",
        "after_box2_sub": "task_id (PK, FK), assignee_email (PK)"
    },
    {
        "badge": "2NF",
        "title": "Second Normal Form",
        "rule": "No partial dependency — non-key columns must depend on the ENTIRE composite PK.",
        "before_headers": ["id", "assignee_email", "project_name"],
        "before_rows": [
            ["1", "tuyen@gmail.com", "Alpha (dup.)"],
            ["1", "huy@gmail.com", "Alpha (dup.)"]
        ],
        "highlight_col": 2,
        "col_widths": [320000, 1080000, 1014016],
        "after_box1_title": "tasks",
        "after_box1_sub": "task_id, project_id (FK)",
        "after_box2_title": "projects",
        "after_box2_sub": "project_id (PK), project_name"
    },
    {
        "badge": "3NF",
        "title": "Third Normal Form",
        "rule": "No transitive dependency — non-key columns must depend directly on the PK.",
        "before_headers": ["id", "title", "st_id", "status_name"],
        "before_rows": [
            ["1", "Fix login", "2", "In Progress"],
            ["2", "Build API", "2", "In Progress"]
        ],
        "highlight_col": 3,
        "col_widths": [300000, 780000, 420000, 914016],
        "after_box1_title": "tasks",
        "after_box1_sub": "task_id, status_id (FK)",
        "after_box2_title": "task_statuses",
        "after_box2_sub": "status_id (PK), name"
    }
]

for col_idx, cdata in enumerate(cols_data):
    col_left = left_margin + col_idx * (col_width + gap)
    inner_left = col_left + 128016
    inner_width = col_width - 256032
    col_center_x = col_left + col_width // 2

    # 1. Column Background Container
    card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_left, card_top, col_width, card_height)
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY_CARD
    card.line.color.rgb = C_SLATE_700
    card.line.width = Pt(1.2)

    # 2. Badge (1NF / 2NF / 3NF)
    badge = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inner_left, card_top + 110000, 470000, 320000)
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_CYAN_DARK
    badge.line.fill.background()
    tf = badge.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = cdata["badge"]
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = C_WHITE

    # 3. Column Title ("First Normal Form", etc.)
    title_box = slide6.shapes.add_textbox(inner_left + 520000, card_top + 115000, inner_width - 520000, 300000)
    tf = title_box.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = cdata["title"]
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = C_WHITE

    # 4. Rule Description
    rule_box = slide6.shapes.add_textbox(inner_left, card_top + 480000, inner_width, 420000)
    tf = rule_box.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = cdata["rule"]
    r.font.size = Pt(8.5)
    r.font.color.rgb = C_SLATE_300

    # 5. "✕ BEFORE — VIOLATION"
    b_lbl = slide6.shapes.add_textbox(inner_left, card_top + 980000, inner_width, 160000)
    tf = b_lbl.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "✕  BEFORE — VIOLATION"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = C_CORAL_RED

    # 6. BEFORE TABLE
    num_rows = len(cdata["before_rows"]) + 1
    num_cols = len(cdata["before_headers"])
    table_top = card_top + 1160000
    table_height = 200000 * num_rows
    tbl_shape = slide6.shapes.add_table(num_rows, num_cols, inner_left, table_top, inner_width, table_height)
    tbl = tbl_shape.table

    # Set Column Widths
    for c_i, w in enumerate(cdata["col_widths"]):
        tbl.columns[c_i].width = w

    # Headers
    for c_i, h_txt in enumerate(cdata["before_headers"]):
        cell = tbl.cell(0, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_RED_HEADER
        cell.margin_top = cell.margin_bottom = cell.margin_left = cell.margin_right = Pt(2)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = h_txt
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = C_RED_TEXT

    # Rows
    for r_i, row_data in enumerate(cdata["before_rows"]):
        for c_i, val in enumerate(row_data):
            cell = tbl.cell(r_i + 1, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_RED_CELL
            cell.margin_top = cell.margin_bottom = cell.margin_left = cell.margin_right = Pt(2)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = val
            r.font.size = Pt(7.5)
            if c_i == cdata["highlight_col"]:
                r.font.bold = True
                r.font.color.rgb = C_CORAL_RED
            else:
                r.font.color.rgb = C_SLATE_200

    # 7. Down Arrow 1
    arrow1_top = card_top + 1820000
    arrow1 = slide6.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, col_center_x - 70000, arrow1_top, 140000, 140000)
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = C_CYAN
    arrow1.line.fill.background()

    # 8. "✓ AFTER — NORMALIZED"
    a_lbl = slide6.shapes.add_textbox(inner_left, card_top + 2000000, inner_width, 160000)
    tf = a_lbl.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "✓  AFTER — NORMALIZED"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = C_MINT_GREEN

    # 9. Box 1 (Top Normalized Box)
    box1_top = card_top + 2190000
    box1_height = 360000
    box1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inner_left, box1_top, inner_width, box1_height)
    box1.fill.solid()
    box1.fill.fore_color.rgb = C_GREEN_CELL
    box1.line.color.rgb = C_MINT_GREEN
    box1.line.width = Pt(1.2)
    tf = box1.text_frame
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(2)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = cdata["after_box1_title"] + "\n"
    r0.font.bold = True
    r0.font.size = Pt(8.5)
    r0.font.color.rgb = C_MINT_GREEN
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = cdata["after_box1_sub"]
    r1.font.size = Pt(8)
    r1.font.color.rgb = C_GREEN_TEXT

    # 10. Down Arrow 2
    arrow2_top = box1_top + box1_height + 40000
    arrow2 = slide6.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, col_center_x - 60000, arrow2_top, 120000, 120000)
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = C_CYAN
    arrow2.line.fill.background()

    # 11. Box 2 (Bottom Normalized Box)
    box2_top = arrow2_top + 160000
    box2_height = 380000
    box2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inner_left, box2_top, inner_width, box2_height)
    box2.fill.solid()
    box2.fill.fore_color.rgb = C_GREEN_CELL
    box2.line.color.rgb = C_MINT_GREEN
    box2.line.width = Pt(1.2)
    tf = box2.text_frame
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(2)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = cdata["after_box2_title"] + "\n"
    r0.font.bold = True
    r0.font.size = Pt(8.5)
    r0.font.color.rgb = C_MINT_GREEN
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = cdata["after_box2_sub"]
    r1.font.size = Pt(8)
    r1.font.color.rgb = C_GREEN_TEXT

# Save file
prs.save(pptx_path)
print(f"Slide 6 perfectly restructured and saved to: {pptx_path}")
