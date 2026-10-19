import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)
slide10 = prs.slides[9] # Slide 10

# Colors
C_WHITE = RGBColor(255, 255, 255)
C_SLATE_100 = RGBColor(241, 245, 249)
C_SLATE_200 = RGBColor(226, 232, 240)
C_SLATE_300 = RGBColor(203, 213, 225)
C_SLATE_400 = RGBColor(148, 163, 184)
C_SLATE_700 = RGBColor(51, 65, 85)
C_SLATE_800 = RGBColor(30, 41, 59)
C_NAVY_DARK = RGBColor(11, 18, 32)
C_TERMINAL_BG = RGBColor(15, 23, 42)

# Accents
C_CYAN = RGBColor(56, 189, 248) # #38BDF8
C_YELLOW_CODE = RGBColor(253, 224, 71) # #FDE047
C_CORAL_RED = RGBColor(248, 113, 113) # #F87171
C_RED_VIBRANT = RGBColor(239, 68, 68)
C_RED_BOX = RGBColor(38, 18, 24)
C_RED_TEXT = RGBColor(254, 202, 202)

C_MINT_GREEN = RGBColor(52, 211, 153) # #34D399
C_GREEN_BOX = RGBColor(8, 38, 30)
C_GREEN_HEADER = RGBColor(6, 78, 59)
C_GREEN_TEXT = RGBColor(167, 243, 208)

# Remove all old shapes except slide number (Shape 0), Title (Shape 1)
shapes_to_remove = []
for idx in range(len(slide10.shapes) - 1, -1, -1):
    shape = slide10.shapes[idx]
    if shape.name not in ['Google Shape;494;p12', 'Google Shape;495;p12']:
        shapes_to_remove.append(shape)

for shape in shapes_to_remove:
    sp = shape._element
    sp.getparent().remove(sp)

print(f"Remaining shapes on Slide 10: {len(slide10.shapes)}")

# Restyle Title
title = slide10.shapes[1]
title.top = 210000
title.left = 365760
for p in title.text_frame.paragraphs:
    for r in p.runs:
        r.font.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = C_WHITE

# Add Subtitle Banner
sub_box = slide10.shapes.add_textbox(365760, 600000, 8412480, 220000)
tf = sub_box.text_frame
tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Comparing PostgreSQL Execution Plans & Performance Live in pgAdmin 4"
r.font.size = Pt(10)
r.font.color.rgb = C_SLATE_400

# Layout Dimensions
col_width = 4050000
col_gap = 312480
left_col = 365760
right_col = left_col + col_width + col_gap
top_base = 880000

# ==============================================================================
# LEFT COLUMN: BEFORE INDEX
# ==============================================================================

# Box 1: Step 1 SQL Query (top: 880000, height: 1100000)
b1 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_col, top_base, col_width, 1100000)
b1.fill.solid()
b1.fill.fore_color.rgb = C_TERMINAL_BG
b1.line.color.rgb = C_SLATE_700
b1.line.width = Pt(1.2)
tf1 = b1.text_frame
tf1.margin_left = Pt(10)
tf1.margin_top = Pt(8)

p0 = tf1.paragraphs[0]
r0 = p0.add_run()
r0.text = "🔴 STEP 1: BEFORE INDEX (Query)"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_CORAL_RED

p1 = tf1.add_paragraph()
r1 = p1.add_run()
r1.text = "EXPLAIN ANALYZE\n"
r1.font.bold = True
r1.font.size = Pt(9)
r1.font.color.rgb = C_CYAN

p2 = tf1.add_paragraph()
r2 = p2.add_run()
r2.text = "SELECT * FROM "
r2.font.size = Pt(9)
r2.font.color.rgb = C_SLATE_200
r2_b = p2.add_run()
r2_b.text = "tasks\n"
r2_b.font.bold = True
r2_b.font.color.rgb = C_WHITE

p3 = tf1.add_paragraph()
r3 = p3.add_run()
r3.text = "WHERE "
r3.font.size = Pt(9)
r3.font.color.rgb = C_SLATE_200
r3_col = p3.add_run()
r3_col.text = "project_id = "
r3_col.font.color.rgb = C_CYAN
r3_val = p3.add_run()
r3_val.text = "'c1284a14-uuid...';"
r3_val.font.color.rgb = C_YELLOW_CODE

# Box 2: Step 1 Expected Output (top: 2060000, height: 1320000)
b2 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_col, 2060000, col_width, 1320000)
b2.fill.solid()
b2.fill.fore_color.rgb = C_RED_BOX
b2.line.color.rgb = C_CORAL_RED
b2.line.width = Pt(1.2)
tf2 = b2.text_frame
tf2.margin_left = Pt(10)
tf2.margin_top = Pt(8)

p0 = tf2.paragraphs[0]
r0 = p0.add_run()
r0.text = "Expected Output: (Sequential Scan)"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_RED_TEXT

p1 = tf2.add_paragraph()
r1 = p1.add_run()
r1.text = "• Seq Scan on tasks  "
r1.font.bold = True
r1.font.size = Pt(8.5)
r1.font.color.rgb = C_CORAL_RED
r1_sub = p1.add_run()
r1_sub.text = "(cost=0.00..250.00 rows=15)"
r1_sub.font.size = Pt(8)
r1_sub.font.color.rgb = C_SLATE_300

p2 = tf2.add_paragraph()
r2 = p2.add_run()
r2.text = "• Filter: (project_id = 'c1284a14...')"
r2.font.size = Pt(8)
r2.font.color.rgb = C_SLATE_400

p3 = tf2.add_paragraph()
r3 = p3.add_run()
r3.text = "⚠ Rows Removed by Filter: 985 (High waste)"
r3.font.bold = True
r3.font.size = Pt(8.5)
r3.font.color.rgb = C_RED_VIBRANT

# Box 3: Step 2 Action Bridge - CREATE INDEX (top: 3460000, height: 1100000)
b3 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_col, 3460000, col_width, 1100000)
b3.fill.solid()
b3.fill.fore_color.rgb = C_TERMINAL_BG
b3.line.color.rgb = C_CYAN
b3.line.width = Pt(1.5)
tf3 = b3.text_frame
tf3.margin_left = Pt(10)
tf3.margin_top = Pt(8)

p0 = tf3.paragraphs[0]
r0 = p0.add_run()
r0.text = "⚡ STEP 2: CREATE THE INDEX"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_CYAN

p1 = tf3.add_paragraph()
r1 = p1.add_run()
r1.text = "CREATE INDEX "
r1.font.bold = True
r1.font.size = Pt(9)
r1.font.color.rgb = C_CYAN
r1_name = p1.add_run()
r1_name.text = "idx_tasks_project_id\n"
r1_name.font.bold = True
r1_name.font.color.rgb = C_WHITE

p2 = tf3.add_paragraph()
r2 = p2.add_run()
r2.text = "ON "
r2.font.size = Pt(9)
r2.font.color.rgb = C_CYAN
r2_target = p2.add_run()
r2_target.text = "tasks(project_id);"
r2_target.font.bold = True
r2_target.font.color.rgb = C_YELLOW_CODE


# ==============================================================================
# RIGHT COLUMN: AFTER INDEX
# ==============================================================================

# Box 4: Step 3 SQL Query (top: 880000, height: 1100000)
b4 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_col, top_base, col_width, 1100000)
b4.fill.solid()
b4.fill.fore_color.rgb = C_TERMINAL_BG
b4.line.color.rgb = C_SLATE_700
b4.line.width = Pt(1.2)
tf4 = b4.text_frame
tf4.margin_left = Pt(10)
tf4.margin_top = Pt(8)

p0 = tf4.paragraphs[0]
r0 = p0.add_run()
r0.text = "🟢 STEP 3: AFTER INDEX (Re-run Query)"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_MINT_GREEN

p1 = tf4.add_paragraph()
r1 = p1.add_run()
r1.text = "EXPLAIN ANALYZE\n"
r1.font.bold = True
r1.font.size = Pt(9)
r1.font.color.rgb = C_CYAN

p2 = tf4.add_paragraph()
r2 = p2.add_run()
r2.text = "SELECT * FROM "
r2.font.size = Pt(9)
r2.font.color.rgb = C_SLATE_200
r2_b = p2.add_run()
r2_b.text = "tasks\n"
r2_b.font.bold = True
r2_b.font.color.rgb = C_WHITE

p3 = tf4.add_paragraph()
r3 = p3.add_run()
r3.text = "WHERE "
r3.font.size = Pt(9)
r3.font.color.rgb = C_SLATE_200
r3_col = p3.add_run()
r3_col.text = "project_id = "
r3_col.font.color.rgb = C_CYAN
r3_val = p3.add_run()
r3_val.text = "'c1284a14-uuid...';"
r3_val.font.color.rgb = C_YELLOW_CODE

# Box 5: Step 3 Expected Output (top: 2060000, height: 1320000)
b5 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_col, 2060000, col_width, 1320000)
b5.fill.solid()
b5.fill.fore_color.rgb = C_GREEN_BOX
b5.line.color.rgb = C_MINT_GREEN
b5.line.width = Pt(1.2)
tf5 = b5.text_frame
tf5.margin_left = Pt(10)
tf5.margin_top = Pt(8)

p0 = tf5.paragraphs[0]
r0 = p0.add_run()
r0.text = "Expected Output: (B-Tree Index Scan)"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_GREEN_TEXT

p1 = tf5.add_paragraph()
r1 = p1.add_run()
r1.text = "• Index Scan using idx_tasks_project_id"
r1.font.bold = True
r1.font.size = Pt(8.5)
r1.font.color.rgb = C_MINT_GREEN

p2 = tf5.add_paragraph()
r2 = p2.add_run()
r2.text = "• Index Cond: (project_id = 'c1284a14...')"
r2.font.size = Pt(8)
r2.font.color.rgb = C_SLATE_300

p3 = tf5.add_paragraph()
r3 = p3.add_run()
r3.text = "✔ Execution: 0.08ms · 0 Rows Removed"
r3.font.bold = True
r3.font.size = Pt(8.5)
r3.font.color.rgb = C_MINT_GREEN

# Box 6: Key Insight Banner (top: 3460000, height: 1100000)
b6 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_col, 3460000, col_width, 1100000)
b6.fill.solid()
b6.fill.fore_color.rgb = C_GREEN_HEADER
b6.line.color.rgb = C_MINT_GREEN
b6.line.width = Pt(1.5)
tf6 = b6.text_frame
tf6.margin_left = Pt(10)
tf6.margin_top = Pt(8)

p0 = tf6.paragraphs[0]
r0 = p0.add_run()
r0.text = "🏆 KEY PERFORMANCE INSIGHT"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_GREEN_TEXT

p1 = tf6.add_paragraph()
r1 = p1.add_run()
r1.text = "• Cost dropped: 250.00 ➔ 8.30 (30x lower)\n"
r1.font.bold = True
r1.font.size = Pt(8.5)
r1.font.color.rgb = C_WHITE

p2 = tf6.add_paragraph()
r2 = p2.add_run()
r2.text = "• Zero wasted rows removed · Resolves in < 1ms"
r2.font.size = Pt(8)
r2.font.color.rgb = C_SLATE_100

prs.save(pptx_path)
print(f"Slide 10 successfully restyled and saved to: {pptx_path}")
