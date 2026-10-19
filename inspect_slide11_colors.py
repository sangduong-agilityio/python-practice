import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

slide11 = prs.slides[10] # Slide 11
print("Slide 11 Shapes Inspection:")
for idx, shape in enumerate(slide11.shapes):
    if shape.has_text_frame:
        print(f"Shape {idx}: '{shape.name}', Text: '{shape.text_frame.text.strip()}'")
