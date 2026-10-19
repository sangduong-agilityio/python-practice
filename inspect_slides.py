import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (3).pptx")

print("Slide width:", prs.slide_width)
print("Slide height:", prs.slide_height)

for idx, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {idx+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            text_str = tf.text.strip().replace('\n', ' | ')
            first_p = tf.paragraphs[0] if tf.paragraphs else None
            font_info = ""
            if first_p and first_p.runs:
                r = first_p.runs[0]
                font_info = f"font={r.font.name}, size={r.font.size}, bold={r.font.bold}"
            print(f"Shape {s_idx} [{shape.name}] (pos: {shape.left},{shape.top},{shape.width},{shape.height}): '{text_str}' {font_info}")
