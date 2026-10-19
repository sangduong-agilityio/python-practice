import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")

for s_idx in range(9, len(prs.slides)):
    slide = prs.slides[s_idx]
    print(f"\n================ SLIDE {s_idx+1} ================")
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace('\n', ' | ')
            if txt.strip():
                print(f"Shape {idx}: '{shape.name}', Text: '{txt[:100]}'")
