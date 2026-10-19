import sys
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (3).pptx"
backup_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (3)_backup.pptx"

# Backup
shutil.copyfile(pptx_path, backup_path)
print(f"Backup created at: {backup_path}")

prs = pptx.Presentation(pptx_path)

def add_header(slide, title_text, subtitle_text):
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(9.0), Inches(0.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.85), Inches(9.0), Inches(0.25))
    tf2 = sub_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle_text.upper()
    p2.font.name = 'Calibri'
    p2.font.size = Pt(9)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(100, 116, 139)

def add_card(slide, left, top, width, height, title, items, bg_rgb=RGBColor(248, 250, 252)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_rgb
    shape.line.color.rgb = RGBColor(203, 213, 225)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_bottom = Inches(0.15)

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.name = 'Calibri'
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(30, 58, 138)
    p0.space_after = Pt(6)

    for item in items:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.name = 'Calibri'
        p.font.size = Pt(9.5)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(4)

# -------------------------------------------------------------
# CREATE SLIDE A: Training Progress Report
# -------------------------------------------------------------
slide_progress = prs.slides.add_slide(prs.slide_layouts[0])
add_header(slide_progress, "Training Roadmap & Progress Status", "Database & System Design Learning Roadmap — Aug 2026")

add_card(slide_progress, Inches(0.45), Inches(1.2), Inches(4.3), Inches(2.2), "✅ COMPLETED PHASES", [
    "Phase 1: Database Design Training Plan (15 Days) — Completed ERD, Normalization 3NF, Constraints, Indexes & Transactions.",
    "Phase 2: System Design Interview by Alex Xu — Completed cover-to-cover all 16 chapters (Web scaling, Rate Limiter, Consistent Hashing, Chat, Drive)."
])

add_card(slide_progress, Inches(4.95), Inches(1.2), Inches(4.3), Inches(2.2), "⏳ IN PROGRESS & EXTENSION REQUEST", [
    "Phase 2.1: Hands-On System Design (Dometrain) — 6/8 chapters done. Requesting 2 additional days to finalize 2 chapters.",
    "Phase 2.2: Python Interview Questions (Dometrain) — Pending. Requesting 1 additional day to complete."
])

add_card(slide_progress, Inches(0.45), Inches(3.55), Inches(8.8), Inches(1.2), "🎯 KEY OUTCOME & PRODUCTION DELIVERY", [
    "Successfully synthesized Database & System Design theory into production code: task-management-api.",
    "Built clean 3-tier REST API with Python 3.13, FastAPI, SQLAlchemy 2.x async, PostgreSQL, Alembic, Pytest & Docker."
])

# Add footer
f_box = slide_progress.shapes.add_textbox(Inches(8.2), Inches(4.8), Inches(1.0), Inches(0.2))
f_tf = f_box.text_frame
f_p = f_tf.paragraphs[0]
f_p.text = "2 / 16"
f_p.font.name = 'Calibri'
f_p.font.size = Pt(10)
f_p.font.color.rgb = RGBColor(100, 116, 139)

# -------------------------------------------------------------
# CREATE SLIDE B: Backend API Architecture & Implementation
# -------------------------------------------------------------
slide_backend = prs.slides.add_slide(prs.slide_layouts[0])
add_header(slide_backend, "Backend REST API Architecture", "Task Management System — Production Implementation")

add_card(slide_backend, Inches(0.45), Inches(1.2), Inches(2.8), Inches(3.5), "🏗️ LAYERED ARCHITECTURE", [
    "Route Handler: FastAPI endpoints with Pydantic v2 validation.",
    "Service Layer: Business rules, ownership checks, custom HTTP exceptions.",
    "Repository Layer: SQLAlchemy 2.x async ORM & raw SQL queries.",
    "Database: PostgreSQL with asyncpg high-performance async driver."
])

add_card(slide_backend, Inches(3.45), Inches(1.2), Inches(2.8), Inches(3.5), "⚡ SYSTEM DESIGN PATTERNS", [
    "3-Tier Architecture: Client ↔ Stateless FastAPI ↔ PostgreSQL.",
    "Read/Write Separation: Tách biệt đường đọc/ghi dữ liệu.",
    "Cache-Aside Pattern: Tối ưu latency đọc dữ liệu công việc.",
    "Async Notifications: Gửi thông báo bất đồng bộ qua Pub/Sub & DLQ."
])

add_card(slide_backend, Inches(6.45), Inches(1.2), Inches(2.8), Inches(3.5), "🔒 SECURITY & QUALITY", [
    "OAuth2 + JWT Auth: Access tokens với hạn 30 phút.",
    "Bcrypt Password Hashing: Mã hóa mật khẩu an toàn.",
    "Pytest Test Suite: Test SQLite in-memory tự động.",
    "Docker Compose & Swagger: Sẵn sàng deploy & Interactive Docs tại /docs."
])

# Add footer
f_box2 = slide_backend.shapes.add_textbox(Inches(8.2), Inches(4.8), Inches(1.0), Inches(0.2))
f_tf2 = f_box2.text_frame
f_p2 = f_tf2.paragraphs[0]
f_p2.text = "15 / 16"
f_p2.font.name = 'Calibri'
f_p2.font.size = Pt(10)
f_p2.font.color.rgb = RGBColor(100, 116, 139)

# Rearrange slides: Move slide_progress to position index 1 (Slide 2), and slide_backend to position index 14 (Slide 15)
# In python-pptx, slide elements in prs.slides._sldIdLst can be reordered
sldIdLst = prs.slides._sldIdLst

# The 2 newly added slides are at the end
progress_elem = sldIdLst[-2]
backend_elem = sldIdLst[-1]

# Remove them from end
sldIdLst.remove(progress_elem)
sldIdLst.remove(backend_elem)

# Insert progress_elem at index 1 (Slide 2)
sldIdLst.insert(1, progress_elem)

# Insert backend_elem at index 14 (before Q&A)
sldIdLst.insert(14, backend_elem)

# Update page footers for all slides
for i, slide in enumerate(prs.slides):
    total = len(prs.slides)
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "/" in txt and ("14" in txt or "16" in txt or txt.replace(" ", "").replace("/", "").isdigit()):
                shape.text_frame.text = f"{i+1} / {total}"
                p = shape.text_frame.paragraphs[0]
                p.font.name = 'Calibri'
                p.font.size = Pt(10)
                p.font.color.rgb = RGBColor(100, 116, 139)

prs.save(pptx_path)
print(f"Presentation successfully updated and saved to: {pptx_path}")
