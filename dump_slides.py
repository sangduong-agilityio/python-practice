import pptx
from pptx.dml.color import RGBColor
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

def dump_slide(s_idx):
    slide = prs.slides[s_idx]
    print(f"\n================ SLIDE {s_idx+1} ================")
    for idx, shape in enumerate(slide.shapes):
        print(f"Shape {idx}: {shape.name} ({shape.shape_type})")
        if shape.has_table:
            t = shape.table
            print(f"  Table: {len(t.rows)}x{len(t.columns)}")
            for r_i, row in enumerate(t.rows):
                for c_i, cell in enumerate(row.cells):
                    fill_type = cell.fill.type if cell.fill else "No fill"
                    print(f"    Cell({r_i},{c_i}): '{cell.text.strip()}'")

dump_slide(5)
dump_slide(6)
