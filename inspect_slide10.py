import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

slide10 = prs.slides[9] # Slide 10
print(f"Slide 10 total shapes: {len(slide10.shapes)}")
for idx, shape in enumerate(slide10.shapes):
    print(f"Shape {idx}: '{shape.name}', type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
    if shape.has_text_frame:
        print(f"    Text: {shape.text_frame.text[:60]}")
