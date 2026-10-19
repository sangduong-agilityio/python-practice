import sys
import pptx
import docx

sys.stdout.reconfigure(encoding='utf-8')

# Read PPTX
pptx_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (4).pptx"
prs = pptx.Presentation(pptx_path)
print(f"=== PPTX: {pptx_path} (Total Slides: {len(prs.slides)}) ===")
for idx, slide in enumerate(prs.slides):
    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip().replace('\n', ' ')
            if t and not title:
                title = t[:70]
    print(f"Slide {idx+1}: {title}")

# Read DOCX
docx_path = r"C:\Users\sangd\Downloads\note demo.docx"
doc = docx.Document(docx_path)
print(f"\n=== DOCX: {docx_path} (Total Paragraphs: {len(doc.paragraphs)}) ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i+1}: {p.text}")
