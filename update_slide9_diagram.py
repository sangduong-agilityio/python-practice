import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
img_path = r"C:\Users\sangd\.gemini\antigravity\brain\27db37f5-881f-470b-8a71-e704d78d7f86\media__1786588882544.png"

prs = pptx.Presentation(pptx_path)
slide9 = prs.slides[8] # 0-indexed 8 = Slide 9

# Colors
C_WHITE = RGBColor(255, 255, 255)
C_SLATE_200 = RGBColor(226, 232, 240)
C_SLATE_300 = RGBColor(203, 213, 225)
C_SLATE_400 = RGBColor(148, 163, 184)
C_SLATE_700 = RGBColor(51, 65, 85)
C_SLATE_800 = RGBColor(30, 41, 59)
C_NAVY_CARD = RGBColor(19, 30, 50)
C_NAVY_DARK = RGBColor(11, 18, 32)
C_CYAN = RGBColor(56, 189, 248)
C_CORAL_RED = RGBColor(248, 113, 113)
C_RED_CELL = RGBColor(45, 20, 26)
C_MINT_GREEN = RGBColor(52, 211, 153)
C_GREEN_CELL = RGBColor(8, 42, 34)

# Keep only slide number and title/subtitle shapes
shapes_to_remove = []
for idx in range(len(slide9.shapes) - 1, -1, -1):
    shape = slide9.shapes[idx]
    if shape.name not in ['Google Shape;438;p11', 'Google Shape;439;p11', 'Google Shape;440;p11']:
        shapes_to_remove.append(shape)

for shape in shapes_to_remove:
    sp = shape._element
    sp.getparent().remove(sp)

print(f"Remaining shapes on Slide 9: {len(slide9.shapes)}")

# Restyle Title and Subtitle
title = slide9.shapes[1]
for p in title.text_frame.paragraphs:
    for r in p.runs:
        r.font.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = C_WHITE

subtitle = slide9.shapes[2]
for p in subtitle.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(11)
        r.font.color.rgb = C_SLATE_400

# Slide width: 9144000, height: 5143500
# Layout:
# 1. Main Central Card Container for Image
card_width = 8412480 # 9.2 inches
card_height = 3200000 # 3.5 inches
card_left = (9144000 - card_width) // 2 # 365760
card_top = 820000

bg_card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_width, card_height)
bg_card.fill.solid()
bg_card.fill.fore_color.rgb = RGBColor(240, 245, 245) # Clean Light Canvas to fit the Crunchy Data image perfectly!
bg_card.line.color.rgb = C_SLATE_700
bg_card.line.width = Pt(1.5)

# Insert Crunchy Data Diagram Image
# Image ratio: 1024 / 682 = 1.501
img_h = 3000000
img_w = int(img_h * 1.501) # ~4503000
img_left = card_left + (card_width - img_w) // 2
img_top = card_top + (card_height - img_h) // 2

pic = slide9.shapes.add_picture(img_path, img_left, img_top, img_w, img_h)

# 2. Bottom Comparison Stats Bars (Side-by-side)
bar_width = 4050000
bar_gap = card_width - 2 * bar_width # 312480
bar_top = card_top + card_height + 140000
bar_height = 540000

# Left Bar: Seq Scan Stats
bar_left = card_left
box_seq = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bar_left, bar_top, bar_width, bar_height)
box_seq.fill.solid()
box_seq.fill.fore_color.rgb = C_RED_CELL
box_seq.line.color.rgb = C_CORAL_RED
box_seq.line.width = Pt(1.2)
tf_seq = box_seq.text_frame
tf_seq.margin_top = Pt(3)
tf_seq.margin_bottom = Pt(2)

p0 = tf_seq.paragraphs[0]
p0.alignment = PP_ALIGN.CENTER
r0 = p0.add_run()
r0.text = "🔴 Sequential Scan (Without Index)\n"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_CORAL_RED

p1 = tf_seq.add_paragraph()
p1.alignment = PP_ALIGN.CENTER
r1 = p1.add_run()
r1.text = "Cost: 0..250.00  ·  Rows Removed: 985 (Slow)  ·  Time: ~12ms"
r1.font.size = Pt(8.5)
r1.font.color.rgb = C_SLATE_200

# Right Bar: Index Scan Stats
bar_right = card_left + bar_width + bar_gap
box_idx = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bar_right, bar_top, bar_width, bar_height)
box_idx.fill.solid()
box_idx.fill.fore_color.rgb = C_GREEN_CELL
box_idx.line.color.rgb = C_MINT_GREEN
box_idx.line.width = Pt(1.2)
tf_idx = box_idx.text_frame
tf_idx.margin_top = Pt(3)
tf_idx.margin_bottom = Pt(2)

p0 = tf_idx.paragraphs[0]
p0.alignment = PP_ALIGN.CENTER
r0 = p0.add_run()
r0.text = "🟢 B-Tree Index Scan (With Index)\n"
r0.font.bold = True
r0.font.size = Pt(9.5)
r0.font.color.rgb = C_MINT_GREEN

p1 = tf_idx.add_paragraph()
p1.alignment = PP_ALIGN.CENTER
r1 = p1.add_run()
r1.text = "Cost: 0.28..8.30  ·  Rows Removed: 0 (Direct)  ·  Time: < 1ms"
r1.font.size = Pt(8.5)
r1.font.color.rgb = C_SLATE_200

# Save presentation
prs.save(pptx_path)
print(f"Slide 9 successfully updated with diagram and saved to: {pptx_path}")
