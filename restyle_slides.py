import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

# Define Palette
C_WHITE = RGBColor(255, 255, 255)
C_SLATE_100 = RGBColor(241, 245, 249)
C_SLATE_200 = RGBColor(226, 232, 240)
C_SLATE_300 = RGBColor(203, 213, 225)
C_SLATE_400 = RGBColor(148, 163, 184)
C_SLATE_500 = RGBColor(100, 116, 139)
C_SLATE_700 = RGBColor(51, 65, 85)
C_SLATE_800 = RGBColor(30, 41, 59)
C_SLATE_900 = RGBColor(15, 23, 42)
C_NAVY_CARD = RGBColor(19, 30, 50)
C_NAVY_DARK = RGBColor(11, 18, 32)

# Accents
C_CYAN = RGBColor(56, 189, 248) # #38BDF8
C_CYAN_DARK = RGBColor(2, 132, 199)
C_CORAL_RED = RGBColor(248, 113, 113) # #F87171
C_RED_VIBRANT = RGBColor(239, 68, 68) # #EF4444
C_RED_HEADER = RGBColor(69, 26, 26) # Deep Red Header
C_RED_CELL = RGBColor(32, 20, 26) # Deep Clean Row
C_RED_TEXT = RGBColor(254, 202, 202) # Light Rose

C_MINT_GREEN = RGBColor(52, 211, 153) # #34D399
C_GREEN_HEADER = RGBColor(6, 78, 59) # Deep Emerald Header
C_GREEN_CELL = RGBColor(8, 38, 32) # Deep Clean Emerald Row
C_GREEN_TEXT = RGBColor(167, 243, 208) # Light Mint

# ==============================================================================
# SLIDE 6 RESTYLE
# ==============================================================================
slide6 = prs.slides[5]

# Style 3 Column Background Containers (Shapes 3, 4, 5)
for s_id in [3, 4, 5]:
    shape = slide6.shapes[s_id]
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_NAVY_CARD
    shape.line.color.rgb = C_SLATE_700
    shape.line.width = Pt(1.2)

# Style 1NF, 2NF, 3NF Badges (Shapes 6, 15, 27)
for s_id in [6, 15, 27]:
    shape = slide6.shapes[s_id]
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_CYAN_DARK
    shape.line.fill.background()
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = C_WHITE

# Style Titles & Texts
for s_id in [7, 16, 28]: # First, Second, Third Normal Form
    shape = slide6.shapes[s_id]
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = C_WHITE

for s_id in [8, 17, 29]: # Descriptions/Rules
    shape = slide6.shapes[s_id]
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = C_SLATE_300

for s_id in [9, 18, 30]: # BEFORE - VIOLATION
    shape = slide6.shapes[s_id]
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = C_CORAL_RED

for s_id in [12, 21, 33]: # AFTER - NORMALIZED
    shape = slide6.shapes[s_id]
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = C_MINT_GREEN

# Down Arrows
for s_id in [11, 20, 24, 32, 36]:
    shape = slide6.shapes[s_id]
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_CYAN
    shape.line.fill.background()

# Helper for Table Styling
def style_table(table, header_bg, header_fg, cell_bg, cell_fg, highlight_cols=None):
    for c_i, cell in enumerate(table.rows[0].cells):
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_bg
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = header_fg
                
    for r_i in range(1, len(table.rows)):
        for c_i, cell in enumerate(table.rows[r_i].cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = cell_bg
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    if highlight_cols and c_i in highlight_cols:
                        r.font.bold = True
                        r.font.color.rgb = C_CORAL_RED
                    else:
                        r.font.color.rgb = cell_fg

# Style Slide 6 Tables
style_table(slide6.shapes[10].table, C_RED_HEADER, C_RED_TEXT, C_RED_CELL, C_SLATE_200, highlight_cols=[2])
style_table(slide6.shapes[13].table, C_GREEN_HEADER, C_GREEN_TEXT, C_GREEN_CELL, C_SLATE_100)

style_table(slide6.shapes[19].table, C_RED_HEADER, C_RED_TEXT, C_RED_CELL, C_SLATE_200, highlight_cols=[2])
# Slide 6 Boxes for 2NF (Shapes 22, 25)
for s_id, t_id in [(22, 23), (25, 26)]:
    box = slide6.shapes[s_id]
    box.fill.solid()
    box.fill.fore_color.rgb = C_GREEN_CELL
    box.line.color.rgb = C_MINT_GREEN
    box.line.width = Pt(1.2)
    txt = slide6.shapes[t_id]
    for p in txt.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = C_GREEN_TEXT

style_table(slide6.shapes[31].table, C_RED_HEADER, C_RED_TEXT, C_RED_CELL, C_SLATE_200, highlight_cols=[3])
# Slide 6 Boxes for 3NF (Shapes 34, 37)
for s_id, t_id in [(34, 35), (37, 38)]:
    box = slide6.shapes[s_id]
    box.fill.solid()
    box.fill.fore_color.rgb = C_GREEN_CELL
    box.line.color.rgb = C_MINT_GREEN
    box.line.width = Pt(1.2)
    txt = slide6.shapes[t_id]
    for p in txt.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = C_GREEN_TEXT

# ==============================================================================
# SLIDE 7 RESTYLE
# ==============================================================================
slide7 = prs.slides[6]

# Main container tasks_flat (Shape 3)
box_tf = slide7.shapes[3]
box_tf.fill.solid()
box_tf.fill.fore_color.rgb = C_NAVY_DARK
box_tf.line.color.rgb = C_SLATE_700
box_tf.line.width = Pt(1.2)

# Title for tasks_flat (Shape 4)
for p in slide7.shapes[4].text_frame.paragraphs:
    for r in p.runs:
        r.font.bold = True
        r.font.color.rgb = C_CYAN

# Table tasks_flat (Shape 5)
table_tf = slide7.shapes[5].table
for c_i, cell in enumerate(table_tf.rows[0].cells):
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_SLATE_800
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = C_CYAN

for r_i in range(1, len(table_tf.rows)):
    row_bg = C_SLATE_900 if r_i % 2 == 1 else C_NAVY_CARD
    for c_i, cell in enumerate(table_tf.rows[r_i].cells):
        cell.fill.solid()
        cell.fill.fore_color.rgb = row_bg
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(8.5)
                # Highlight duplicates in Red
                if r_i in [1, 2, 3] and c_i in [2, 3]:
                    r.font.color.rgb = C_CORAL_RED
                    r.font.bold = True
                else:
                    r.font.color.rgb = C_SLATE_200

# 4 Disaster / Anomaly Cards (Shapes: Card, Dot, Title, Desc)
cards_data = [
    (6, 7, 8, 9),    # Data Redundancy
    (10, 11, 12, 13), # Update Anomaly
    (14, 15, 16, 17), # Insert Anomaly
    (18, 19, 20, 21)  # Delete Anomaly
]

for card_id, dot_id, title_id, desc_id in cards_data:
    # Card Box
    card = slide7.shapes[card_id]
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY_CARD
    card.line.color.rgb = C_SLATE_700
    card.line.width = Pt(1.2)
    
    # Red Dot
    dot = slide7.shapes[dot_id]
    dot.fill.solid()
    dot.fill.fore_color.rgb = C_RED_VIBRANT
    dot.line.fill.background()
    
    # Title
    t_shape = slide7.shapes[title_id]
    for p in t_shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = C_CORAL_RED
            
    # Description
    d_shape = slide7.shapes[desc_id]
    for p in d_shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(9.5)
            r.font.color.rgb = C_SLATE_300

# Resolution Banner at Bottom (Shapes 22, 23)
res_box = slide7.shapes[22]
res_box.fill.solid()
res_box.fill.fore_color.rgb = C_GREEN_CELL
res_box.line.color.rgb = C_MINT_GREEN
res_box.line.width = Pt(1.5)

res_txt = slide7.shapes[23]
for p in res_txt.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = C_GREEN_TEXT

# Save file
output_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs.save(output_path)
print(f"Successfully styled and saved to: {output_path}")
