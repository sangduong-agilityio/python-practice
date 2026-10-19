import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)
slide7 = prs.slides[6]

print(f"Slide 7 total shapes: {len(slide7.shapes)}")
for idx, shape in enumerate(slide7.shapes):
    print(f"Shape {idx}: '{shape.name}', type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
