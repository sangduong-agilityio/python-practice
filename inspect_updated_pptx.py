import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")

for s_idx in [5, 6]: # Slide 6 and 7 (0-indexed 5 and 6)
    slide = prs.slides[s_idx]
    print(f"\n================ SLIDE {s_idx+1} ================")
    for idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        print(f"Shape {idx}: type={shape_type}, name='{shape.name}', left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace('\n', ' \\n ')
            print(f"    Text: {txt[:120]}")
        if shape.has_table:
            print(f"    Table: rows={len(shape.table.rows)}, cols={len(shape.table.columns)}")
            for r_idx, row in enumerate(shape.table.rows):
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                print(f"      Row {r_idx}: {row_txt}")
