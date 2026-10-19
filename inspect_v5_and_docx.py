import pptx
import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\sangd\Downloads\[SangDuong] Database Training - 2026 (5).pptx"
docx_path = r"C:\Users\sangd\Downloads\note demo (1).docx"

prs = pptx.Presentation(pptx_path)
print(f"=== PPTX: {pptx_path} ===")
print(f"Total slides: {len(prs.slides)}")

for s_idx, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {s_idx+1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip().replace('\n', ' | ')
            if txt:
                print(f"  [{shape.name}]: {txt[:90]}")

doc = docx.Document(docx_path)
print(f"\n=== DOCX: {docx_path} ===")
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")
for i, p in enumerate(doc.paragraphs[:25]):
    if p.text.strip():
        print(f"  P{i}: {p.text[:90]}")
