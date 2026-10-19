import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r"C:\Users\sangd\Downloads\note demo (1).docx"
doc = docx.Document(docx_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("Slide ") or "—" in txt or txt.startswith("###") or txt.startswith("=="):
        print(f"P{idx}: {txt}")
