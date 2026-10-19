import pptx
from pptx.util import Inches, Pt
import sys
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

img_path = r"C:\Users\sangd\.gemini\antigravity\brain\27db37f5-881f-470b-8a71-e704d78d7f86\media__1786588882544.png"
with Image.open(img_path) as im:
    print(f"Image format={im.format}, size={im.size}, mode={im.mode}")

pptx_path = r"C:\Users\sangd\Downloads\_SangDuong__Database_Training_-_2026__4__updated.pptx"
prs = pptx.Presentation(pptx_path)

# Let's inspect slide 9/10 ("Indexing: Seq Scan vs Index Scan")
for idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame and "Seq Scan vs Index Scan" in shape.text_frame.text:
            print(f"Found 'Seq Scan vs Index Scan' on Slide {idx+1}")
